"""Build the SUTRA keyword-monitor Excel report from a scan JSON + curated summaries.

Usage: .venv/bin/python build_report.py data/scan_2026-10.json reports/<name>.xlsx "October 2026 (month to date)"
"""
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, ".")
from sutra_scan import KEYWORDS  # noqa: E402

# Curated per-measure analysis (written after reading each document), keyed by measure number.
ANALYSIS = json.load(open("data/analysis.json"))
STRONG_TEXT_KW = {
    "DDEC",
    "Departamento de Desarrollo Económico y Comercio",
    "Department of Economic Development and Commerce",
    "Invest Puerto Rico",
    "Ley 60",
    "código de incentivos",
}

FONT = "Arial"
NAVY = "1F3864"
HDR_FILL = PatternFill("solid", start_color=NAVY)
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
REL_FILL = {
    "High": PatternFill("solid", start_color="F8CBAD"),
    "Medium": PatternFill("solid", start_color="FFE699"),
    "Low-Medium": PatternFill("solid", start_color="FFF2CC"),
    "Low": PatternFill("solid", start_color="E2EFDA"),
}


def last_event(m):
    if not m["eventos"]:
        return ""
    e = m["eventos"][-1]
    return f"{e['evento']} ({e['fecha']})"


def write_table(ws, headers, widths, rows, link_cols=()):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
            if c in link_cols and v:
                cell.hyperlink = v
                cell.font = LINK
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def detail_rows(measures, kw_field):
    order = {"High": 0, "Medium": 1, "Low-Medium": 2, "Low": 3}
    measures = sorted(measures, key=lambda m: (order[ANALYSIS[m["numero"]]["relevance"]], m["fecha_radicacion"]))
    rows = []
    for m in measures:
        a = ANALYSIS[m["numero"]]
        kws = m[kw_field] if kw_field == "keywords" else [k for k in m["keywords_texto"] if k in STRONG_TEXT_KW]
        rows.append([
            m["numero"], a["relevance"], m["tipo"], m["fecha_radicacion"], ", ".join(m["autores"]),
            a["committee"], last_event(m), ", ".join(kws), m["titulo"], a["summary"], a["why"],
            m["url"], m["documentos"][0] if m["documentos"] else "",
        ])
    return rows


DETAIL_HEADERS = [
    "Measure", "Relevance to InvestPR", "Type", "Filed", "Author(s)", "Committee", "Latest status",
    "Keywords matched", "Official title (Spanish)", "Quick summary", "Why it matters", "SUTRA link", "Document link",
]
DETAIL_WIDTHS = [11, 12, 20, 11, 20, 22, 24, 22, 55, 60, 45, 30, 30]


def main(scan_path, out_path, label):
    data = json.load(open(scan_path))
    pr = [m for m in data if m["is_proyecto_o_resolucion"]]
    title_hits = [m for m in pr if m["keywords"]]
    text_hits = [m for m in pr if not m["keywords"] and STRONG_TEXT_KW & set(m.get("keywords_texto", []))]
    missing = [m["numero"] for m in title_hits + text_hits if m["numero"] not in ANALYSIS]
    if missing:
        sys.exit(f"Add ANALYSIS entries for: {missing}")
    period = f"{min(m['fecha_radicacion'] for m in data)} to {max(m['fecha_radicacion'] for m in data)}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 95
    ws["A1"] = f"SUTRA Legislative Keyword Monitor — {label}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    info = [
        ("Source", "Sistema Único de Trámite Legislativo (SUTRA), https://sutra.oslpr.org"),
        ("Filing period scanned", period),
        ("Report generated", date.today().isoformat()),
        ("Measures filed in period (all types)", len(data)),
        ("Proyectos and resoluciones in scope", len(pr)),
        ("Keyword matches in title", len(title_hits)),
        ("Additional matches in document body", len(text_hits)),
        ("Keywords tracked", "; ".join(KEYWORDS)),
        ("Method",
         "Every measure filed in the period was read from SUTRA. Only proyectos and resoluciones are in scope. "
         "Tier 1 = keyword appears in the official title (per Stakeholder Affairs guidance). "
         "Tier 2 = keyword does not appear in the title but the filed document text mentions DDEC, the Department's "
         "full name, Ley 60 / Código de Incentivos, or Invest Puerto Rico. Matching ignores accents and capitalization; "
         "'Ley 60' matches 'Ley 60-2019' and 'Ley Núm. 60' but not 'Ley 600'."),
        ("Caveats",
         f"{sum(1 for m in pr if not m['documentos'])} in-scope measures had no document posted yet, so only their titles were checked. "
         "Generic body-text mentions of 'desarrollo económico' or 'inversión' were frequent and mostly incidental, so they are "
         "listed only on the 'All Measures' tab. Summaries and relevance ratings are AI-generated from the filed text and "
         "should be reviewed before sharing externally."),
    ]
    for r, (k, v) in enumerate(info, 3):
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = BODY, Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")

    for title, rows in (("Title Matches", detail_rows(title_hits, "keywords")),
                        ("Body Text Mentions", detail_rows(text_hits, "keywords_texto"))):
        ws = wb.create_sheet(title)
        write_table(ws, DETAIL_HEADERS, DETAIL_WIDTHS, rows, link_cols=(12, 13))
        for r in range(2, len(rows) + 2):
            ws.cell(row=r, column=2).fill = REL_FILL[ws.cell(row=r, column=2).value]

    ws = wb.create_sheet("All Measures")
    rows = []
    for m in sorted(data, key=lambda m: (m["fecha_radicacion"], m["numero"])):
        rows.append([
            m["numero"], m["tipo"], m["fecha_radicacion"], "Yes" if m["is_proyecto_o_resolucion"] else "No",
            ", ".join(m["keywords"]), ", ".join(m.get("keywords_texto", [])), m["titulo"],
            ", ".join(m["autores"]), last_event(m), m["url"], m["documentos"][0] if m["documentos"] else "",
        ])
    write_table(
        ws,
        ["Measure", "Type", "Filed", "In scope", "Keywords in title", "Keywords in document text", "Official title (Spanish)",
         "Author(s)", "Latest status", "SUTRA link", "Document link"],
        [11, 26, 11, 8, 24, 30, 70, 22, 26, 30, 30], rows, link_cols=(10, 11),
    )

    wb.save(out_path)
    print(f"saved {out_path}: {len(title_hits)} title matches, {len(text_hits)} body-text mentions, {len(data)} total")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
